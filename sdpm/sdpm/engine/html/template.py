# Copyright Amazon.com, Inc. or its affiliates. All Rights Reserved.
# SPDX-License-Identifier: MIT-0
"""Convert a .pptx template into the HTML-mode artifacts.

``theme.css``      canvas size, pt scale, ground, fonts, theme colours, and per
                   layout the placeholder boxes and text styles
                   (``section[data-layout="…"] :is([data-ph="0"],[data-ph="title"])``)
``template.json``  layouts, placeholders (data-ph keys, boxes, resolved styles),
                   chrome (slide number / footer), decorations, grounds
``layouts/*.svg``  decoration-only renders (LibreOffice + poppler; optional)

The structure (theme.css + template.json) needs python-pptx only. Results are
cached per template content hash; the bundled templates ship pre-converted
artifacts under ``sdpm/templates/html/<stem>/``.
"""

from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from sdpm.engine import slide_size_px
from sdpm.engine.html import decorations as _decorations
from sdpm.engine.html.textstyle import load_theme, placeholder_type, resolve_placeholder_style

CONVERTER_VERSION = "1"
TEMPLATE_JSON_VERSION = 1

_CHROME_TYPES = ("dt", "ftr", "sldNum")
_ALIASES = {"title": ["title"], "ctrTitle": ["title"], "subTitle": ["subtitle"]}
_FILE_RE = re.compile(r"[^A-Za-z0-9_-]+")
_JUSTIFY = {"top": "flex-start", "middle": "center", "bottom": "flex-end"}
_AUTONUM = (("arabic", "decimal"), ("alphaLc", "lower-alpha"), ("alphaUc", "upper-alpha"),
            ("romanLc", "lower-roman"), ("romanUc", "upper-roman"))

# Process-local memo of failed renders (cache dir → status) so a broken
# LibreOffice conversion is attempted once per process, not on every build.
_FAILED_RENDERS: dict[str, dict] = {}


@dataclass
class TemplateHtml:
    """Converted template: template.json data, theme.css text, and the artifact paths."""

    info: dict
    theme_css: str
    dir: Path
    svg_paths: dict[str, Path] = field(default_factory=dict)
    ground_svg_paths: dict[str, Path] = field(default_factory=dict)
    decorations_status: dict | None = None


def cache_key(template_path: str | Path) -> str:
    """Content hash of the template plus the converter version (not a security hash)."""
    data = Path(template_path).read_bytes()
    digest = hashlib.sha256(data + CONVERTER_VERSION.encode(), usedforsecurity=False)
    return digest.hexdigest()[:16]


def layout_file_name(index: int, name: str) -> str:
    """``NN-<sanitised name>`` with the 1-based layout position (unique per template)."""
    return f"{index:02d}-{_FILE_RE.sub('_', name)}"


def template_json_text(info: dict) -> str:
    """Deterministic serialisation used for both the cache and the bundled artifacts."""
    return json.dumps(info, ensure_ascii=False, indent=2, sort_keys=True) + "\n"


# ── structure (python-pptx only) ─────────────────────────────────────────

def _placeholder_text(ph) -> str:
    if not getattr(ph, "has_text_frame", False):
        return ""
    return ph.text_frame.text.replace("\x0b", " / ").replace("\v", " / ").strip()


def _ph_bbox(ph, emu: float) -> dict[str, int] | None:
    if ph.left is None or ph.top is None or ph.width is None or ph.height is None:
        return None
    return {"x": round(ph.left / emu), "y": round(ph.top / emu),
            "w": round(ph.width / emu), "h": round(ph.height / emu)}


def _layout_entry(prs, layout, index: int, theme: dict, canvas: tuple[int, int],
                  curated: dict[str, str | None]) -> dict[str, Any]:
    emu = theme["emuPerPx"]
    name = layout.name or f"layout-{index:02d}"
    entry: dict[str, Any] = {"name": name, "file": layout_file_name(index, name),
                             "curated": name in curated}
    if curated.get(name):
        entry["notes"] = curated[name]

    placeholders = []
    chrome: dict[str, Any] = {}
    for ph in layout.placeholders:
        el = ph._element
        ph_type = placeholder_type(el)
        bbox = _ph_bbox(ph, emu)
        if bbox is None:
            continue
        style = resolve_placeholder_style(prs, layout, el, theme)
        if ph_type in _ALIASES and ph_type != "subTitle":
            # Titles are single-paragraph; deeper levels only echo defaultTextStyle.
            style = {**style, "levels": style["levels"][:1]}
        text = _placeholder_text(ph)
        if ph_type in _CHROME_TYPES:
            chrome_text = {**style, "levels": style["levels"][:1]}
            if ph_type == "sldNum":
                chrome["slideNumber"] = {"bbox": bbox, "text": chrome_text}
            elif ph_type == "ftr" and text:
                chrome["footer"] = {"bbox": bbox, "text": chrome_text, "value": text}
            continue
        item: dict[str, Any] = {"key": str(ph.placeholder_format.idx), "type": ph_type,
                                "bbox": bbox, "text": style}
        if ph_type in _ALIASES:
            item["aliases"] = list(_ALIASES[ph_type])
        if text:
            item["prompt"] = text
        placeholders.append(item)
    entry["placeholders"] = placeholders
    if chrome:
        entry["chrome"] = chrome
    decorations = _decorations.layout_decorations(prs, layout, *canvas)
    if decorations:
        entry["decorations"] = decorations
    ground = _decorations.layout_ground(prs, layout, theme)
    entry["ground"] = ground
    if ground["kind"] == "baked":
        entry["groundSvg"] = True
    return entry


def build_template_info(template_path: str | Path) -> dict[str, Any]:
    """template.json data for a .pptx (python-pptx only; no LibreOffice)."""
    from pptx import Presentation

    from sdpm.engine.analyzer import analyze_template, extract_fonts, extract_theme_colors

    template_path = Path(template_path)
    prs = Presentation(str(template_path))
    theme = load_theme(template_path, prs)
    width, height = slide_size_px(prs.slide_width, prs.slide_height)
    pt_per_px = round((int(prs.slide_width) / 12700) / 1920, 6)
    analysis = analyze_template(template_path)
    curated = {item["name"]: item.get("notes") for item in analysis.get("layouts", [])}
    master = prs.slide_masters[0]
    master_bg = _decorations.master_ground(master, theme)
    side_fonts = extract_fonts(template_path)
    layouts = [
        _layout_entry(prs, layout, i, theme, (width, height), curated)
        for i, layout in enumerate(master.slide_layouts, 1)
    ]
    return {
        "version": TEMPLATE_JSON_VERSION,
        "sourceHash": cache_key(template_path),
        "converterVersion": CONVERTER_VERSION,
        "template": template_path.stem,
        "slideSize": {"width": width, "height": height, "ptPerPx": pt_per_px},
        "ptScale": round(0.75 / pt_per_px, 4),
        "fonts": {
            "major": theme["fonts"]["major"],
            "minor": theme["fonts"]["minor"],
            "halfwidth": side_fonts.get("halfwidth"),
            "fullwidth": side_fonts.get("fullwidth"),
        },
        "themeColors": extract_theme_colors(template_path),
        "ground": master_bg.get("color") if master_bg["kind"] == "solid" else None,
        "layouts": layouts,
    }


# ── theme.css ────────────────────────────────────────────────────────────

def _num(value: float) -> str:
    return format(round(float(value), 4), "g")


def css_string(value: str) -> str:
    """Double-quoted CSS string with backslash, quote and newline escaped."""
    escaped = value.replace("\\", "\\\\").replace('"', '\\"').replace("\n", "\\a ")
    return f'"{escaped}"'


def _font_css(level: dict) -> str | None:
    ref = level.get("fontRef")
    if ref in ("major", "minor"):
        return f"var(--sdpm-font-{ref})"
    if level.get("font"):
        return css_string(level["font"])
    return None


def _pt(value: float) -> str:
    return f"calc({_num(value)}pt * var(--sdpm-pt-scale))"


def _spacing_css(value: dict) -> str:
    if "pt" in value:
        return _pt(value["pt"])
    return f"{_num(value['pct'] * 1.2)}em"


def _line_height_css(value: dict) -> str:
    if "pt" in value:
        return _pt(value["pt"])
    return _num(value["pct"] * 1.2)


def _text_decls(level: dict, *, with_font: bool) -> list[str]:
    decls = []
    if with_font:
        font = _font_css(level)
        if font:
            decls.append(f"font-family:var(--font-family, {font}), {font}")
    if "sizePt" in level:
        decls.append(f"font-size:{_pt(level['sizePt'])}")
    if level.get("color"):
        decls.append(f"color:{level['color']}")
    if level.get("align"):
        decls.append(f"text-align:{level['align']}")
    if "bold" in level:
        decls.append(f"font-weight:{700 if level['bold'] else 400}")
    if "italic" in level:
        decls.append(f"font-style:{'italic' if level['italic'] else 'normal'}")
    if "caps" in level:
        decls.append(f"text-transform:{'uppercase' if level['caps'] else 'none'}")
    return decls


def _list_style(bullet: dict | None) -> str | None:
    if not bullet:
        return None
    if bullet["kind"] == "none":
        return "none"
    if bullet["kind"] == "char":
        return css_string(bullet.get("char", "•") + " ")
    scheme = bullet.get("scheme", "")
    for prefix, css in _AUTONUM:
        if scheme.startswith(prefix):
            return css
    return "decimal"


def _placeholder_rules(selector: str, ph: dict) -> list[str]:
    text = ph["text"]
    levels = text["levels"]
    first = levels[0] if levels else {}
    bbox, insets = ph["bbox"], text["insets"]
    box = [
        "position:absolute",
        f"left:{bbox['x']}px", f"top:{bbox['y']}px", f"width:{bbox['w']}px", f"height:{bbox['h']}px",
        "box-sizing:border-box", "margin:0",
        f"padding:{_num(insets['top'])}px {_num(insets['right'])}px "
        f"{_num(insets['bottom'])}px {_num(insets['left'])}px",
        "display:flex", "flex-direction:column",
        f"justify-content:{_JUSTIFY.get(text['anchor'], 'flex-start')}",
    ]
    box += _text_decls(first, with_font=True)
    box.append(f"line-height:{_line_height_css(first['lineSpacing']) if 'lineSpacing' in first else '1.2'}")
    rules = [f"{selector}{{{';'.join(box)}}}",
             f"{selector} :is(p,h1,h2,h3,h4,h5,h6,ul,ol,li){{margin:0}}",
             f"{selector} :is(ul,ol){{padding:0}}"]
    para = []
    if "spaceBefore" in first:
        para.append(f"margin-top:{_spacing_css(first['spaceBefore'])}")
    if "spaceAfter" in first:
        para.append(f"margin-bottom:{_spacing_css(first['spaceAfter'])}")
    if para:
        rules.append(f"{selector} :is(p,li){{{';'.join(para)}}}")

    prev_margin: float = 0
    prev_decls: list[str] = _text_decls(first, with_font=True)
    prev_indent = 0
    prev_list = None
    prev_line = first.get("lineSpacing")
    for level in levels:
        n = level["level"]
        sel = selector + " li" * n
        decls = []
        margin = level.get("marginLeftPx", prev_margin)
        if margin != prev_margin:
            decls.append(f"padding-left:{_num(margin - prev_margin)}px")
        prev_margin = margin
        indent = level.get("indentPx", prev_indent)
        if indent != prev_indent:
            decls.append(f"text-indent:{_num(indent)}px")
        prev_indent = indent
        list_style = _list_style(level.get("bullet")) or prev_list
        if n == 1:
            decls.append("list-style-position:inside")
        if list_style and list_style != prev_list:
            decls.append(f"list-style-type:{list_style}")
        prev_list = list_style
        if n > 1:
            # Nested <li> inherit from their parent <li>: emit only what changes.
            level_decls = _text_decls(level, with_font=True)
            decls += [d for d in level_decls if d not in prev_decls]
            prev_decls = level_decls or prev_decls
            for key, prop in (("spaceBefore", "margin-top"), ("spaceAfter", "margin-bottom")):
                if key in level and level[key] != first.get(key):
                    decls.append(f"{prop}:{_spacing_css(level[key])}")
            line = level.get("lineSpacing", prev_line)
            if line != prev_line and line is not None:
                decls.append(f"line-height:{_line_height_css(line)}")
            prev_line = line
        if decls:
            rules.append(f"{sel}{{{';'.join(decls)}}}")
        bullet = level.get("bullet") or {}
        marker = []
        if bullet.get("color"):
            marker.append(f"color:{bullet['color']}")
        if bullet.get("sizePct"):
            marker.append(f"font-size:{_num(bullet['sizePct'])}%")
        if bullet.get("font"):
            marker.append(f"font-family:{css_string(bullet['font'])}")
        if marker:
            rules.append(f"{sel}::marker{{{';'.join(marker)}}}")
    return rules


def build_theme_css(info: dict) -> str:
    """theme.css text for template.json data (deterministic)."""
    size = info["slideSize"]
    fonts = info["fonts"]
    root = [
        f"--sdpm-canvas-w:{size['width']}px",
        f"--sdpm-canvas-h:{size['height']}px",
        f"--sdpm-pt-scale:{_num(info['ptScale'])}",
        f"--sdpm-template-bg:{info.get('ground') or '#FFFFFF'}",
        f"--sdpm-font-major:{css_string(fonts.get('major') or 'sans-serif')}",
        f"--sdpm-font-minor:{css_string(fonts.get('minor') or 'sans-serif')}",
    ]
    for role, color in sorted(info.get("themeColors", {}).items()):
        root.append(f"--sdpm-theme-{role}:{color}")
    lines = [
        "/* Generated by sdpm.engine.html.template — do not edit. */",
        f":root{{{';'.join(root)}}}",
    ]
    for layout in info["layouts"]:
        section = f"section[data-layout={css_string(layout['name'])}]"
        ground = layout.get("ground", {})
        if ground.get("kind") == "solid" and ground.get("color"):
            lines.append(f"{section}{{--sdpm-template-bg:{ground['color']}}}")
        for ph in layout["placeholders"]:
            keys = [ph["key"], *ph.get("aliases", [])]
            selector = f"{section} :is({','.join(f'[data-ph={css_string(k)}]' for k in keys)})"
            lines.extend(_placeholder_rules(selector, ph))
    return "\n".join(lines) + "\n"


# ── cache / bundled lookup ───────────────────────────────────────────────

def _read_info(path: Path) -> dict | None:
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None
    return data if isinstance(data, dict) else None


def _write_atomic(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(f".{path.name}.{os.getpid()}.tmp")
    try:
        tmp.write_text(text, encoding="utf-8", newline="\n")
        os.replace(tmp, path)
    except BaseException:
        tmp.unlink(missing_ok=True)
        raise


def _render_entries(info: dict) -> list[dict]:
    return [{"name": layout["name"], "file": layout["file"], "index": i,
             "baked": bool(layout.get("groundSvg"))}
            for i, layout in enumerate(info["layouts"])]


def _svgs_complete(info: dict, layouts_dir: Path) -> bool:
    for layout in info["layouts"]:
        if not (layouts_dir / f"{layout['file']}.svg").is_file():
            return False
        if layout.get("groundSvg") and not (layouts_dir / f"{layout['file']}.ground.svg").is_file():
            return False
    return True


def _result(directory: Path, info: dict, theme_css: str, status: dict | None, stem: str) -> TemplateHtml:
    layouts_dir = directory / "layouts"
    svgs: dict[str, Path] = {}
    grounds: dict[str, Path] = {}
    for layout in info["layouts"]:
        svg = layouts_dir / f"{layout['file']}.svg"
        if svg.is_file():
            svgs[layout["name"]] = svg
        ground = layouts_dir / f"{layout['file']}.ground.svg"
        if ground.is_file():
            grounds[layout["name"]] = ground
    return TemplateHtml(info={**info, "template": stem}, theme_css=theme_css, dir=directory,
                        svg_paths=svgs, ground_svg_paths=grounds, decorations_status=status)


def _render_into(template_path: Path, info: dict, cache_dir: Path) -> dict | None:
    """Render decoration SVGs into ``cache_dir/layouts``; returns the failure status or None."""
    memo = _FAILED_RENDERS.get(str(cache_dir))
    if memo is not None:
        return memo
    staging = Path(tempfile.mkdtemp(prefix=".render-", dir=cache_dir))
    try:
        rendered = _decorations.render_decoration_svgs(
            template_path, _render_entries(info), staging / "layouts", work_dir=staging / "work")
        if rendered["status"] is not None:
            if rendered["status"].get("missing"):
                return rendered["status"]  # tools vanished: report, do not memoise
            _FAILED_RENDERS[str(cache_dir)] = rendered["status"]
            return rendered["status"]
        layouts_dir = cache_dir / "layouts"
        layouts_dir.mkdir(parents=True, exist_ok=True)
        for svg in sorted((staging / "layouts").glob("*.svg")):
            os.replace(svg, layouts_dir / svg.name)
        return None
    finally:
        shutil.rmtree(staging, ignore_errors=True)


def load_template_html(
    template_path: str | Path,
    *,
    decorations: bool = True,
    cache_root: str | Path | None = None,
    bundled_root: str | Path | None = None,
) -> TemplateHtml:
    """Converted HTML artifacts for a .pptx template (bundled, cached, or freshly built).

    Bundled artifacts (``bundled_root/<stem>/``) are used when their template.json
    ``sourceHash`` / ``converterVersion`` match the template; otherwise the cache
    ``cache_root/<cache_key>/`` is filled — structure first (python-pptx only), then
    the decoration SVGs when ``decorations`` is true and LibreOffice + poppler exist.
    Missing tools or a failed render come back as ``decorations_status``; never raises
    for them.
    """
    from sdpm.config import CACHE_DIR, TEMPLATES_DIR

    template_path = Path(template_path)
    stem = template_path.stem
    key = cache_key(template_path)

    bundled_dir = Path(bundled_root if bundled_root is not None else TEMPLATES_DIR / "html") / stem
    info = _read_info(bundled_dir / "template.json")
    css_path = bundled_dir / "theme.css"
    if (info and info.get("sourceHash") == key and info.get("converterVersion") == CONVERTER_VERSION
            and css_path.is_file()):
        return _result(bundled_dir, info, css_path.read_text(encoding="utf-8"), None, stem)

    cache_dir = Path(cache_root if cache_root is not None else CACHE_DIR / "template-html") / key
    info = _read_info(cache_dir / "template.json")
    css_path = cache_dir / "theme.css"
    theme_css = css_path.read_text(encoding="utf-8") if css_path.is_file() else ""
    if info is None or not theme_css or info.get("converterVersion") != CONVERTER_VERSION:
        info = build_template_info(template_path)
        theme_css = build_theme_css(info)
        _write_atomic(cache_dir / "theme.css", theme_css)
        _write_atomic(cache_dir / "template.json", template_json_text(info))

    status: dict | None = None
    if decorations and not _svgs_complete(info, cache_dir / "layouts"):
        status = _decorations.decoration_tools_status()
        if status is None:
            status = _render_into(template_path, info, cache_dir)
    return _result(cache_dir, info, theme_css, status, stem)


__all__ = [
    "CONVERTER_VERSION",
    "TemplateHtml",
    "build_template_info",
    "build_theme_css",
    "cache_key",
    "layout_file_name",
    "load_template_html",
    "template_json_text",
]

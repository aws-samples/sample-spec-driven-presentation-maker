# Copyright Amazon.com, Inc. or its affiliates. All Rights Reserved.
# SPDX-License-Identifier: MIT-0
"""Template decorations for HTML decks: bounding boxes, layout grounds, SVG renders.

Decorations are the non-placeholder shapes a layout shows (its own plus the
master's unless ``showMasterSp="0"``): logos, copyright lines, bands. Their
boxes come from the XML (python-pptx only). The pixels come from LibreOffice
(PPTX → PDF) and poppler (``pdftocairo -svg``): one page per layout with every
slide placeholder removed and a sentinel ground colour that is stripped from
the SVG afterwards, because the HTML deck paints the ground in CSS.
"""

from __future__ import annotations

import math
import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any

from sdpm.engine import emu_per_px
from sdpm.engine.converter.color import _resolve_scheme_color
from sdpm.engine.converter.constants import _NS
from sdpm.engine.html.textstyle import resolve_color
from sdpm.engine.preview.environment import preview_environment, soffice_path

SENTINEL_RGB = (0x01, 0x02, 0x03)
_SENTINEL_HEX = "#010203"
_PDFTOCAIRO_TIMEOUT = 60

DECORATIONS_NOTE = (
    "Template structure converted; only the template decorations (logos, bands) are "
    "missing from the HTML deck. Install the missing tools to include them."
)
CONVERSION_FAILED_NOTE = (
    "LibreOffice could not render the template decorations; the HTML deck is built "
    "without them (placeholders and styles are unaffected)."
)

_SHAPE_TAGS = ("sp", "pic", "grpSp", "graphicFrame", "cxnSp")


def _local(tag) -> str:
    return tag.rsplit("}", 1)[-1] if isinstance(tag, str) else ""


def _xfrm_of(el):
    tag = _local(el.tag)
    if tag == "grpSp":
        return el.find("p:grpSpPr/a:xfrm", _NS)
    if tag == "graphicFrame":
        return el.find("p:xfrm", _NS)
    return el.find("p:spPr/a:xfrm", _NS)


def _is_placeholder(el) -> bool:
    return el.find("./*/p:nvPr/p:ph", _NS) is not None


def _shape_name(el) -> str:
    c_nv_pr = el.find("./*/p:cNvPr", _NS)
    return c_nv_pr.get("name", "") if c_nv_pr is not None else ""


def _rotated_bbox(x: float, y: float, w: float, h: float, rot_deg: float) -> tuple[float, float, float, float]:
    if not rot_deg % 180:
        return x, y, w, h
    rad = math.radians(rot_deg)
    cos, sin = abs(math.cos(rad)), abs(math.sin(rad))
    nw, nh = w * cos + h * sin, w * sin + h * cos
    cx, cy = x + w / 2, y + h / 2
    return cx - nw / 2, cy - nh / 2, nw, nh


def _bbox_px(el, emu: float) -> dict[str, int] | None:
    xfrm = _xfrm_of(el)
    if xfrm is None:
        return None
    off = xfrm.find("a:off", _NS)
    ext = xfrm.find("a:ext", _NS)
    if off is None or ext is None:
        return None
    x, y = int(off.get("x", 0)), int(off.get("y", 0))
    w, h = int(ext.get("cx", 0)), int(ext.get("cy", 0))
    rot = int(xfrm.get("rot", 0)) / 60000
    x, y, w, h = _rotated_bbox(x, y, w, h, rot)
    return {"x": round(x / emu), "y": round(y / emu), "w": round(w / emu), "h": round(h / emu)}


def _decoration_shapes(container):
    tree = container._element.find("p:cSld/p:spTree", _NS)
    if tree is None:
        return []
    return [el for el in tree if _local(el.tag) in _SHAPE_TAGS and not _is_placeholder(el)]


def shows_master_shapes(layout) -> bool:
    return layout._element.get("showMasterSp") not in ("0", "false")


def placeholder_bboxes(layout, emu: float) -> list[dict[str, int]]:
    """Pixel boxes of the layout's placeholders (inherited positions resolved)."""
    boxes = []
    for ph in layout.placeholders:
        if ph.left is None or ph.width is None:
            continue
        boxes.append({"x": round(ph.left / emu), "y": round(ph.top / emu),
                      "w": round(ph.width / emu), "h": round(ph.height / emu)})
    return boxes


def _contains(outer: dict, inner: dict) -> bool:
    return (outer["x"] <= inner["x"] and outer["y"] <= inner["y"]
            and outer["x"] + outer["w"] >= inner["x"] + inner["w"]
            and outer["y"] + outer["h"] >= inner["y"] + inner["h"])


def layout_decorations(prs, layout, canvas_w: int, canvas_h: int) -> list[dict[str, Any]]:
    """Decoration boxes a layout shows, in drawing order (master shapes first).

    ``kind`` is ``ground`` (covers >= 50% of the canvas), ``band`` (fully contains
    a placeholder of this layout) or ``obstacle``. Groups count as one box.
    """
    emu = emu_per_px(prs.slide_width)
    placeholders = placeholder_bboxes(layout, emu)
    sources = []
    if shows_master_shapes(layout):
        sources.append(("master", layout.slide_master))
    sources.append(("layout", layout))
    canvas_area = canvas_w * canvas_h
    out = []
    for source, container in sources:
        for el in _decoration_shapes(container):
            bbox = _bbox_px(el, emu)
            if bbox is None or bbox["w"] <= 0 or bbox["h"] <= 0:
                continue
            if bbox["w"] * bbox["h"] >= 0.5 * canvas_area:
                kind = "ground"
            elif any(_contains(bbox, ph) for ph in placeholders):
                kind = "band"
            else:
                kind = "obstacle"
            out.append({"name": _shape_name(el), "source": source, "bbox": bbox, "kind": kind})
    return out


def _bg_of(container):
    return container._element.find("p:cSld/p:bg", _NS)


def _fill_kind(fill_el) -> str:
    return _local(fill_el.tag) if fill_el is not None else ""


def background_ground(bg, theme: dict) -> dict[str, str]:
    """``{"kind": "solid", "color"}`` or ``{"kind": "baked"}`` for a ``p:bg`` element (None = theme bg1)."""
    if bg is not None:
        bg_pr = bg.find("p:bgPr", _NS)
        if bg_pr is not None:
            fill = next((c for c in bg_pr if _local(c.tag).endswith("Fill")), None)
            kind = _fill_kind(fill)
            if kind == "solidFill":
                color = resolve_color(fill, theme)
                if color:
                    return {"kind": "solid", "color": color}
            elif kind in ("blipFill", "gradFill", "pattFill", "grpFill"):
                return {"kind": "baked"}
        bg_ref = bg.find("p:bgRef", _NS)
        if bg_ref is not None:
            idx = int(bg_ref.get("idx", "0"))
            ref_color = resolve_color(bg_ref, theme)
            styles = theme["bgFillStyles"] if idx >= 1001 else theme["fillStyles"]
            pos = idx - 1001 if idx >= 1001 else idx - 1
            entry = styles[pos] if 0 <= pos < len(styles) else None
            kind = _fill_kind(entry)
            if kind in ("blipFill", "gradFill", "pattFill"):
                return {"kind": "baked"}
            if kind == "solidFill":
                scheme = entry.find("a:schemeClr", _NS)
                if scheme is not None and scheme.get("val") == "phClr" and ref_color:
                    return {"kind": "solid", "color": ref_color}
                color = resolve_color(entry, theme)
                if color:
                    return {"kind": "solid", "color": color}
            if ref_color:
                return {"kind": "solid", "color": ref_color}
    bg1 = _resolve_scheme_color("bg1", theme["colors"], theme["mapping"]) or "#FFFFFF"
    return {"kind": "solid", "color": bg1}


def master_ground(master, theme: dict) -> dict[str, str]:
    return background_ground(_bg_of(master), theme)


def layout_ground(prs, layout, theme: dict) -> dict[str, str]:
    """Effective ground of a layout: its own ``p:bg``, else the master's."""
    bg = _bg_of(layout)
    if bg is None:
        bg = _bg_of(layout.slide_master)
    return background_ground(bg, theme)


# ── SVG render (LibreOffice + poppler) ──────────────────────────────────

def decoration_tools_status() -> dict | None:
    """None when LibreOffice and pdftocairo exist; else the unavailable status (cheap check)."""
    missing = []
    if soffice_path() is None:
        missing.append("libreoffice")
    if shutil.which("pdftocairo") is None:
        missing.append("poppler")
    if not missing:
        return None
    return unavailable_status(missing)


def unavailable_status(missing: list[str], note: str = DECORATIONS_NOTE) -> dict:
    install = preview_environment()["install"] or "install LibreOffice and poppler (pdftocairo) and put them on PATH"
    return {"status": "unavailable", "missing": missing, "install": install, "note": note}


def _layout_by_entry(master, entry: dict):
    layouts = list(master.slide_layouts)
    index = entry.get("index")
    if isinstance(index, int) and 0 <= index < len(layouts):
        return layouts[index]
    match = None
    for i, layout in enumerate(layouts):
        if (layout.name or f"layout-{i + 1:02d}") == entry["name"]:
            match = layout
    return match


def build_decoration_deck(template_path: str | Path, layouts: list[dict], out_path: Path) -> list[tuple[str, dict]]:
    """Write the temp pptx: one sentinel page per layout, then one plain page per baked layout.

    Returns the page plan ``[("svg" | "ground", layout_entry), ...]`` in page order.
    """
    from pptx import Presentation
    from pptx.dml.color import RGBColor

    prs = Presentation(str(template_path))
    sld_id_lst = prs.slides._sldIdLst
    for sld_id in list(sld_id_lst):
        prs.part.drop_rel(sld_id.rId)
        sld_id_lst.remove(sld_id)
    master = prs.slide_masters[0]

    plan: list[tuple[str, dict]] = []

    def add_page(entry: dict, sentinel: bool) -> None:
        layout = _layout_by_entry(master, entry)
        if layout is None:
            return
        slide = prs.slides.add_slide(layout)
        for ph in list(slide.placeholders):
            ph._element.getparent().remove(ph._element)
        if sentinel:
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = RGBColor(*SENTINEL_RGB)
        plan.append(("svg" if sentinel else "ground", entry))

    for entry in layouts:
        add_page(entry, sentinel=True)
    for entry in layouts:
        if entry.get("baked"):
            add_page(entry, sentinel=False)
    prs.save(str(out_path))
    return plan


_FILL_RE = re.compile(r'''(?:\bfill\s*=\s*"([^"]+)"|\bfill\s*:\s*([^;"]+))''')
_SHAPE_RE = re.compile(r"<(path|rect)\b([^>]*?)(?:/>|>\s*</\1\s*>)", re.S)
_RGB_PCT_RE = re.compile(r"rgb\(\s*([\d.]+)%\s*,\s*([\d.]+)%\s*,\s*([\d.]+)%\s*\)")
_RGB_INT_RE = re.compile(r"rgb\(\s*(\d+)\s*,\s*(\d+)\s*,\s*(\d+)\s*\)")


def _is_sentinel_fill(value: str) -> bool:
    value = value.strip().lower()
    if value == _SENTINEL_HEX:
        return True
    channels: tuple[float, ...] | None = None
    m = _RGB_PCT_RE.fullmatch(value)
    if m:
        channels = tuple(float(c) * 255 / 100 for c in m.groups())
    else:
        m = _RGB_INT_RE.fullmatch(value)
        if m:
            channels = tuple(float(c) for c in m.groups())
    if channels is None:
        return False
    return all(abs(c - s) <= 1.0 for c, s in zip(channels, SENTINEL_RGB))


_VIEWBOX_RE = re.compile(r'<svg\b[^>]*\bviewBox="([^"]+)"', re.S)
_D_RE = re.compile(r'\bd="([^"]*)"')
_NUM_RE = re.compile(r"-?\d+(?:\.\d+)?(?:[eE]-?\d+)?")


def _covers_page(attrs: str, page_w: float, page_h: float) -> bool:
    m = _D_RE.search(attrs)
    if m is None:
        return False
    nums = [float(n) for n in _NUM_RE.findall(m.group(1))]
    xs, ys = nums[0::2], nums[1::2]
    if not xs or not ys:
        return False
    tol = 2.0
    return (min(xs) <= tol and min(ys) <= tol
            and max(xs) >= page_w - tol and max(ys) >= page_h - tol)


def _fill_of(attrs: str) -> str | None:
    m = _FILL_RE.search(attrs)
    return (m.group(1) or m.group(2)) if m else None


def strip_sentinel(svg_text: str) -> str:
    """Remove the page ground: every ``<path>`` / ``<rect>`` in the sentinel colour,
    plus the full-page underlay the PDF export paints before it."""
    matches = list(_SHAPE_RE.finditer(svg_text))
    sentinel = [m for m in matches if _is_sentinel_fill(_fill_of(m.group(2)) or "")]
    if not sentinel:
        return svg_text
    first = sentinel[0].start()
    defs_end = svg_text.find("</defs>")
    vb = _VIEWBOX_RE.search(svg_text)
    page = [float(v) for v in vb.group(1).split()] if vb else []
    drop = {m.start() for m in sentinel}
    if len(page) == 4:
        for m in matches:
            if defs_end < m.start() < first and _fill_of(m.group(2)) and _covers_page(m.group(2), page[2], page[3]):
                drop.add(m.start())
    out, pos = [], 0
    for m in matches:
        if m.start() in drop:
            out.append(svg_text[pos:m.start()])
            pos = m.end()
    out.append(svg_text[pos:])
    return "".join(out)


def _pdf_page_to_svg(pdftocairo: str, pdf: Path, page: int, out: Path) -> bool:
    try:
        subprocess.run(  # nosec B603 # nosemgrep: python.lang.security.audit.dangerous-subprocess-use-audit
            [pdftocairo, "-svg", "-f", str(page), "-l", str(page), str(pdf), str(out)],
            capture_output=True, timeout=_PDFTOCAIRO_TIMEOUT, check=True, stdin=subprocess.DEVNULL,
        )
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired, FileNotFoundError, OSError):
        return False
    return out.is_file()


def render_decoration_svgs(
    template_path: str | Path,
    layouts: list[dict],
    out_dir: Path,
    *,
    work_dir: Path,
) -> dict[str, Any]:
    """Render decoration SVGs for ``layouts`` (``{name, file, baked, index?}``) into ``out_dir``.

    Writes ``<file>.svg`` (sentinel ground stripped) per layout and
    ``<file>.ground.svg`` (full render) per baked-ground layout. Returns
    ``{"svgs": {name: Path}, "ground_svgs": {name: Path}, "status": None | dict}``;
    never raises — missing tools or a failed conversion come back as ``status``.
    """
    result: dict[str, Any] = {"svgs": {}, "ground_svgs": {}, "status": None}
    status = decoration_tools_status()
    soffice = soffice_path()
    pdftocairo = shutil.which("pdftocairo")
    if status is not None or soffice is None or pdftocairo is None:
        result["status"] = status or unavailable_status(["libreoffice", "poppler"])
        return result

    from sdpm.engine.preview.backend import LibreOfficeBackend

    failed = {"status": "unavailable", "missing": [], "install": "", "note": CONVERSION_FAILED_NOTE}
    work_dir.mkdir(parents=True, exist_ok=True)
    out_dir.mkdir(parents=True, exist_ok=True)
    tmp = Path(tempfile.mkdtemp(prefix="sdpm-decor-", dir=work_dir))
    try:
        deck = tmp / "decorations.pptx"
        try:
            plan = build_decoration_deck(template_path, layouts, deck)
        except Exception:  # noqa: BLE001 — a broken template must not break the deck build
            result["status"] = failed
            return result
        pdf = tmp / "decorations.pdf"
        if not plan or not LibreOfficeBackend(soffice).export_pdf(deck, pdf, work_dir=tmp):
            result["status"] = failed
            return result
        for page, (kind, entry) in enumerate(plan, 1):
            raw = tmp / f"page-{page}.svg"
            if not _pdf_page_to_svg(pdftocairo, pdf, page, raw):
                result["status"] = failed
                return result
            text = raw.read_text(encoding="utf-8")
            if kind == "svg":
                target = out_dir / f"{entry['file']}.svg"
                target.write_text(strip_sentinel(text), encoding="utf-8")
                result["svgs"][entry["name"]] = target
            else:
                target = out_dir / f"{entry['file']}.ground.svg"
                target.write_text(text, encoding="utf-8")
                result["ground_svgs"][entry["name"]] = target
        return result
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

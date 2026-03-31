# Copyright Amazon.com, Inc. or its affiliates. All Rights Reserved.
# SPDX-License-Identifier: MIT-0
"""Table element."""
from pptx.dml.color import RGBColor
from pptx.util import Emu, Pt


class TableMixin:
    """Mixin providing table element methods."""

    def _add_table(self, slide, elem):
        from pptx.enum.text import PP_ALIGN
        from lxml import etree
    
        headers = elem.get("headers", [])
        rows = elem.get("rows", [])
        cols = len(headers) if headers else (len(rows[0]) if rows else 0)
        row_count = len(rows) + (1 if headers else 0)
    
        if cols == 0 or row_count == 0:
            return
    
        x = self._px_to_emu(elem.get("x", 77))
        y = self._px_to_emu(elem.get("y", 270))
        width = self._px_to_emu(elem.get("width", 1766))
        height = self._px_to_emu(elem.get("height")) if elem.get("height") else Emu(row_count * 400000)
    
        tbl_shape = slide.shapes.add_table(row_count, cols, x, y, width, height)
        table = tbl_shape.table
        nsmap = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
    
        # Column widths
        col_widths = elem.get("colWidths")
        if col_widths:
            for i, w in enumerate(col_widths):
                if i < len(table.columns):
                    table.columns[i].width = self._px_to_emu(w)
    
        # Row heights
        row_heights = elem.get("rowHeights")
        if row_heights:
            for i, h in enumerate(row_heights):
                if i < len(table.rows):
                    table.rows[i].height = self._px_to_emu(h)
    
        # Table style properties
        tbl_pr = table._tbl.find('a:tblPr', nsmap)
        if tbl_pr is not None:
            for attr in ['firstRow', 'lastRow', 'firstCol', 'lastCol', 'bandRow', 'bandCol']:
                if elem.get(attr):
                    tbl_pr.set(attr, '1')
                elif attr in tbl_pr.attrib:
                    del tbl_pr.attrib[attr]
            # Resolve tableStyle name → GUID
            table_style = elem.get("tableStyle", "")
            style_id = None
            if table_style:
                style_id = self._table_style_map.get(table_style)
                if not style_id:
                    import sys
                    print(f"Warning: Unknown tableStyle '{table_style}', using template default", file=sys.stderr)
            if not style_id:
                # No style specified or unknown — use template default
                style_id = self._default_table_style_id
            if style_id:
                existing = tbl_pr.find('a:tableStyleId', nsmap)
                if existing is not None:
                    existing.text = style_id
                else:
                    el = etree.SubElement(tbl_pr, f'{{{nsmap["a"]}}}tableStyleId')
                    el.text = style_id
            else:
                # No template styles at all — remove tableStyleId, use theme-based fallback
                existing = tbl_pr.find('a:tableStyleId', nsmap)
                if existing is not None:
                    tbl_pr.remove(existing)
    
        # Merge cells first (scan all rows for gridSpan/rowSpan)
        all_rows = [headers] + rows if headers else rows
        for ri, row_data in enumerate(all_rows):
            for ci, cell_val in enumerate(row_data):
                if not isinstance(cell_val, dict):
                    continue
                gs = cell_val.get("gridSpan", 1)
                rs = cell_val.get("rowSpan", 1)
                if gs > 1 or rs > 1:
                    try:
                        table.cell(ri, ci).merge(table.cell(ri + rs - 1, ci + gs - 1))
                    except Exception:
                        pass
    
        # Helper to apply cell properties
        def apply_cell(cell, cell_val, fallback_fill=None, fallback_text_color=None):
            text = cell_val if isinstance(cell_val, str) else cell_val.get("text", "")
            is_dict = isinstance(cell_val, dict)
    
            # Skip merged-away cells
            if is_dict and cell_val.get("merged"):
                return
    
            # Fill
            if is_dict and "gradient" in cell_val:
                g = cell_val["gradient"]
                from sdpm.builder.formatting import _build_grad_fill_element
                tcPr = cell._tc.get_or_add_tcPr()
                grad_fill = _build_grad_fill_element(g)
                tcPr.append(grad_fill)
            elif is_dict and "fill" in cell_val:
                f = cell_val["fill"]
                if f and f != "none":
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor.from_string(f.lstrip('#'))
                elif f == "none":
                    cell.fill.background()
            elif fallback_fill:
                cell.fill.solid()
                cell.fill.fore_color.rgb = fallback_fill
    
            # Vertical alignment (default: middle)
            _va_table_map = {"top": "t", "middle": "ctr", "bottom": "b"}
            va = cell_val.get("verticalAlign", "middle") if is_dict else "middle"
            anchor_val = _va_table_map.get(va, "ctr")
            tc_pr = cell._tc.find(f'{{{nsmap["a"]}}}tcPr')
            if tc_pr is None:
                tc_pr = etree.SubElement(cell._tc, f'{{{nsmap["a"]}}}tcPr')
            tc_pr.set('anchor', anchor_val)
    
            # Margins
            if is_dict and cell_val.get("margins"):
                tc_pr = cell._tc.find(f'{{{nsmap["a"]}}}tcPr')
                if tc_pr is None:
                    tc_pr = etree.SubElement(cell._tc, f'{{{nsmap["a"]}}}tcPr')
                m = cell_val["margins"]
                for side, attr in [('left', 'marL'), ('right', 'marR'), ('top', 'marT'), ('bottom', 'marB')]:
                    if side in m:
                        tc_pr.set(attr, str(self._px_to_emu(m[side])))
    
            # Borders
            if is_dict and cell_val.get("borders"):
                tc_pr = cell._tc.find(f'{{{nsmap["a"]}}}tcPr')
                if tc_pr is None:
                    tc_pr = etree.SubElement(cell._tc, f'{{{nsmap["a"]}}}tcPr')
                tag_map = {"left": "lnL", "right": "lnR", "top": "lnT", "bottom": "lnB"}
                for side, bdr in cell_val["borders"].items():
                    tag = tag_map.get(side)
                    if not tag:
                        continue
                    # Remove existing
                    existing = tc_pr.find(f'{{{nsmap["a"]}}}{tag}')
                    if existing is not None:
                        tc_pr.remove(existing)
                    ln = etree.SubElement(tc_pr, f'{{{nsmap["a"]}}}{tag}')
                    if bdr.get("width"):
                        ln.set('w', str(int(bdr["width"] * 12700)))
                    if bdr.get("fill") == "none":
                        etree.SubElement(ln, f'{{{nsmap["a"]}}}noFill')
                    elif bdr.get("color"):
                        sf = etree.SubElement(ln, f'{{{nsmap["a"]}}}solidFill')
                        srgb = etree.SubElement(sf, f'{{{nsmap["a"]}}}srgbClr')
                        srgb.set('val', bdr["color"].lstrip('#'))
    
            # Text
            para = cell.text_frame.paragraphs[0]
            text_color = fallback_text_color
            if is_dict:
                if cell_val.get("fontColor"):
                    text_color = RGBColor.from_string(cell_val["fontColor"].lstrip('#'))
                # Alignment
                align = cell_val.get("align")
                if align:
                    para.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}.get(align)
                # Font size
                font_size = cell_val.get("fontSize")
            else:
                font_size = None
    
            self._apply_styled_text(para, str(text), default_color=text_color, no_default_color=text_color is None)
    
            # Apply bold/italic/fontSize on all runs
            if is_dict:
                for run in para.runs:
                    if cell_val.get("bold"):
                        run.font.bold = True
                    if cell_val.get("italic"):
                        run.font.italic = True
                    if cell_val.get("underline"):
                        run.font.underline = True
                    if font_size:
                        run.font.size = Pt(font_size)
    
        # Fill headers
        has_table_style = bool(style_id)
        # Theme-aware fallback colors for unstyled tables (derived from theme)
        text_hex = self.theme_colors["text"].lstrip("#")
        bg_hex = self.theme_colors["background"].lstrip("#")
        header_fill = RGBColor.from_string(text_hex)
        header_text = RGBColor.from_string(bg_hex)
        row_bg = RGBColor.from_string(bg_hex)
        # Slightly shift background for alternating rows
        br, bg, bb = int(bg_hex[:2], 16), int(bg_hex[2:4], 16), int(bg_hex[4:6], 16)
        shift = -12 if (0.299 * br + 0.587 * bg + 0.114 * bb) > 128 else 12
        alt = f"{max(0,min(255,br+shift)):02X}{max(0,min(255,bg+shift)):02X}{max(0,min(255,bb+shift)):02X}"
        row_bg_alt = RGBColor.from_string(alt)
        row_text = RGBColor.from_string(text_hex)
        if headers:
            for ci, hdr in enumerate(headers):
                is_roundtrip_cell = isinstance(hdr, dict) and "borders" in hdr
                apply_cell(table.cell(0, ci), hdr,
                           fallback_fill=None if (has_table_style or is_roundtrip_cell) else header_fill,
                           fallback_text_color=None if (has_table_style or is_roundtrip_cell) else header_text)

        # Fill data rows
        start_row = 1 if headers else 0
        for ri, row in enumerate(rows):
            bg = row_bg_alt if ri % 2 == 1 else row_bg
            for ci, val in enumerate(row):
                is_roundtrip_cell = isinstance(val, dict) and "borders" in val
                apply_cell(table.cell(start_row + ri, ci), val,
                           fallback_fill=None if (has_table_style or is_roundtrip_cell) else bg,
                           fallback_text_color=None if (has_table_style or is_roundtrip_cell) else row_text)
    


# Copyright Amazon.com, Inc. or its affiliates. All Rights Reserved.
# SPDX-License-Identifier: MIT-0
"""Tests for the HTML-mode template conversion (sdpm.engine.html)."""

from __future__ import annotations

import json
import re
import shutil
from pathlib import Path

import pytest
from pptx import Presentation

from conftest import DECORATED_BAND_RGB, DECORATED_LOGO_PX
from sdpm.config import TEMPLATES_DIR
from sdpm.engine.html import decorations, template
from sdpm.engine.html.decorations import (
    layout_decorations,
    layout_ground,
    strip_sentinel,
)
from sdpm.engine.html.template import (
    build_template_info,
    build_theme_css,
    cache_key,
    load_template_html,
    template_json_text,
)
from sdpm.engine.html.textstyle import load_theme, resolve_placeholder_style
from sdpm.engine.preview.environment import soffice_path

_BUNDLED = ("blank-light", "blank-dark")
_needs_lo = pytest.mark.skipif(
    soffice_path() is None or shutil.which("pdftocairo") is None,
    reason="LibreOffice and poppler (pdftocairo) required",
)


def _layout(prs, name):
    return next(layout for layout in prs.slide_masters[0].slide_layouts if layout.name == name)


def _ph(layout, idx):
    return next(ph for ph in layout.placeholders if ph.placeholder_format.idx == idx)


def _no_bundle(tmp_path: Path) -> Path:
    return tmp_path / "no-bundle"


# ── text styles ──────────────────────────────────────────────────────────

def test_textstyle_resolves_blank_light_title_and_body_levels() -> None:
    path = TEMPLATES_DIR / "blank-light.pptx"
    prs = Presentation(str(path))
    theme = load_theme(path, prs)
    layout = _layout(prs, "Title and Content")

    title = resolve_placeholder_style(prs, layout, _ph(layout, 0)._element, theme)
    assert title["anchor"] == "middle"
    assert title["insets"] == {"left": 14.4, "top": 7.2, "right": 14.4, "bottom": 7.2}
    lvl1 = title["levels"][0]
    assert (lvl1["sizePt"], lvl1["color"], lvl1["align"]) == (44.0, "#000000", "center")
    assert lvl1["fontRef"] == "major" and lvl1["font"] == "Calibri"
    assert lvl1["bullet"] == {"kind": "none"}

    body = resolve_placeholder_style(prs, layout, _ph(layout, 1)._element, theme)
    assert body["anchor"] == "top"
    levels = body["levels"]
    assert [lvl["sizePt"] for lvl in levels] == [32.0, 28.0, 24.0, 20.0, 20.0]
    assert levels[0]["bullet"] == {"kind": "char", "char": "•", "font": "Arial"}
    assert levels[1]["bullet"]["char"] == "–"
    assert levels[0]["marginLeftPx"] == 54 and levels[0]["indentPx"] == -54
    assert levels[0]["spaceBefore"] == {"pct": 0.2}
    assert levels[0]["fontRef"] == "minor"


def test_textstyle_layout_lststyle_overrides_master() -> None:
    path = TEMPLATES_DIR / "blank-light.pptx"
    prs = Presentation(str(path))
    theme = load_theme(path, prs)
    layout = _layout(prs, "Title Slide")
    title = resolve_placeholder_style(prs, layout, _ph(layout, 0)._element, theme)
    assert title["levels"][0]["align"] == "left"  # layout lvl1pPr algn="l" beats titleStyle ctr
    subtitle = resolve_placeholder_style(prs, layout, _ph(layout, 1)._element, theme)
    first = subtitle["levels"][0]
    assert first["bullet"] == {"kind": "none"}
    assert first["color"] != "#000000"  # tx1 with tint 75%
    assert first["sizePt"] == 32.0  # inherited from master bodyStyle


# ── decorations and grounds ──────────────────────────────────────────────

def test_layout_decorations_kinds_for_fixture(decorated_template: Path) -> None:
    prs = Presentation(str(decorated_template))
    items = {d["name"]: d for d in layout_decorations(prs, _layout(prs, "Title Only"), 1920, 1080)}
    assert set(items) == {"Logo", "Copyright", "Header Band"}
    x, y, w, h = DECORATED_LOGO_PX
    assert items["Logo"] == {"name": "Logo", "source": "master", "kind": "obstacle",
                             "bbox": {"x": x, "y": y, "w": w, "h": h}}
    assert items["Copyright"]["kind"] == "obstacle"
    assert items["Header Band"]["kind"] == "band"
    assert items["Header Band"]["source"] == "layout"
    assert items["Header Band"]["bbox"]["w"] == 1920


def test_layout_decorations_hide_master_shapes_when_show_master_sp_is_off(decorated_template: Path) -> None:
    prs = Presentation(str(decorated_template))
    assert layout_decorations(prs, _layout(prs, "Blank"), 1920, 1080) == []
    names = {d["name"] for d in layout_decorations(prs, _layout(prs, "Title and Content"), 1920, 1080)}
    assert names == {"Logo", "Copyright"}


def test_layout_ground_solid_and_baked(decorated_template: Path) -> None:
    light = TEMPLATES_DIR / "blank-light.pptx"
    prs = Presentation(str(light))
    assert layout_ground(prs, _layout(prs, "Title Only"), load_theme(light, prs)) == {
        "kind": "solid", "color": "#FFFFFF"}
    dark = TEMPLATES_DIR / "blank-dark.pptx"
    prs = Presentation(str(dark))
    ground = layout_ground(prs, _layout(prs, "Title Only"), load_theme(dark, prs))
    assert ground["kind"] == "solid" and ground["color"] != "#FFFFFF"

    prs = Presentation(str(decorated_template))
    theme = load_theme(decorated_template, prs)
    assert layout_ground(prs, _layout(prs, "Section Header"), theme) == {"kind": "baked"}


def test_strip_sentinel_removes_ground_and_page_underlay() -> None:
    svg = (
        '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 960 540"><defs>'
        '<clipPath id="c"><path d="M 0 0 L 960 0 L 960 540 L 0 540 Z"/></clipPath></defs>'
        '<g clip-path="url(#c)">'
        '<path fill="rgb(100%, 100%, 100%)" d="M 0 0 L 959 0 L 959 540 L 0 540 Z"/>'
        '<path fill="rgb(0.392151%, 0.784302%, 1.176453%)" d="M 0 0 L 959 0 L 959 540 L 0 540 Z"/>'
        '<rect fill="#010203" x="0" y="0" width="10" height="10"/>'
        '<path fill="rgb(100%, 60%, 0%)" d="M 0 0 L 960 0 L 960 80 L 0 80 Z"/>'
        '<path fill="rgb(100%, 100%, 100%)" d="M 0 0 L 960 0 L 960 540 L 0 540 Z"/>'
        "</g></svg>"
    )
    out = strip_sentinel(svg)
    assert "0.392151%" not in out and "#010203" not in out
    assert out.count('fill="rgb(100%, 100%, 100%)"') == 1  # underlay gone, later full-page shape kept
    assert "rgb(100%, 60%, 0%)" in out
    assert "<clipPath" in out


# ── template.json / theme.css ────────────────────────────────────────────

def test_template_info_shape_for_fixture(decorated_template: Path) -> None:
    info = build_template_info(decorated_template)
    assert info["version"] == 1
    assert info["sourceHash"] == cache_key(decorated_template)
    assert info["converterVersion"] == template.CONVERTER_VERSION
    assert info["slideSize"] == {"width": 1920, "height": 1080, "ptPerPx": 0.5}
    assert info["ptScale"] == 1.5
    assert set(info["fonts"]) == {"major", "minor", "halfwidth", "fullwidth"}
    assert info["ground"] == "#FFFFFF"
    assert info["themeColors"]["accent1"].startswith("#")

    layouts = {layout["name"]: layout for layout in info["layouts"]}
    assert len(layouts) == 11
    assert layouts["Title Slide"]["curated"] is True and layouts["Title Only"]["curated"] is True
    assert layouts["Blank"]["curated"] is False
    assert layouts["Title Slide"]["file"] == "01-Title_Slide"

    title_slide = {ph["key"]: ph for ph in layouts["Title Slide"]["placeholders"]}
    assert set(title_slide) == {"0", "1"}  # dt / ftr / sldNum are chrome, not placeholders
    assert title_slide["0"]["aliases"] == ["title"] and title_slide["0"]["type"] == "ctrTitle"
    assert title_slide["1"]["aliases"] == ["subtitle"]
    assert title_slide["0"]["prompt"] == "Title"
    assert set(title_slide["0"]["bbox"]) == {"x", "y", "w", "h"}
    assert "slideNumber" in layouts["Title Slide"]["chrome"]
    assert "footer" not in layouts["Title Slide"]["chrome"]  # no literal footer text

    kinds = {d["name"]: d["kind"] for d in layouts["Title Only"]["decorations"]}
    assert kinds == {"Logo": "obstacle", "Copyright": "obstacle", "Header Band": "band"}
    assert "decorations" not in layouts["Blank"]
    assert layouts["Section Header"]["ground"] == {"kind": "baked"}
    assert layouts["Section Header"]["groundSvg"] is True
    assert layouts["Title Only"]["ground"] == {"kind": "solid", "color": "#FFFFFF"}


def test_theme_css_shape() -> None:
    info = build_template_info(TEMPLATES_DIR / "blank-light.pptx")
    css = build_theme_css(info)
    assert ":root{--sdpm-canvas-w:1920px;--sdpm-canvas-h:1080px;--sdpm-pt-scale:1.5;" in css
    assert "--sdpm-template-bg:#FFFFFF" in css
    assert '--sdpm-font-major:"Calibri"' in css
    assert "--sdpm-theme-accent1:" in css
    assert 'section[data-layout="Title Only"]{--sdpm-template-bg:#FFFFFF}' in css
    selector = ':where(section[data-layout="Title and Content"]) :is([data-ph="0"],[data-ph="title"])'
    assert f"{selector}{{position:absolute;" in css
    assert "font-size:calc(44pt * var(--sdpm-pt-scale))" in css
    assert "font-family:var(--font-family, var(--sdpm-font-major)), var(--sdpm-font-major)" in css
    assert ':where(section[data-layout="Title and Content"]) :is([data-ph="1"]) li li{' in css
    assert 'list-style-type:"• "' in css


def test_theme_css_escapes_layout_names() -> None:
    info = build_template_info(TEMPLATES_DIR / "blank-light.pptx")
    info["layouts"][0]["name"] = 'A "quoted" \\ name'
    css = build_theme_css(info)
    assert 'section[data-layout="A \\"quoted\\" \\\\ name"]' in css


def test_theme_css_pt_scale_for_4x3(template_4x3: Path) -> None:
    info = build_template_info(template_4x3)
    assert info["ptScale"] == 2.0
    assert info["slideSize"]["height"] == 1440
    assert "--sdpm-pt-scale:2;" in build_theme_css(info)


# ── cache ────────────────────────────────────────────────────────────────

def _fake_render(calls: list):
    def render(template_path, layouts, out_dir, *, work_dir):
        calls.append(template_path)
        out_dir.mkdir(parents=True, exist_ok=True)
        svgs, grounds = {}, {}
        for entry in layouts:
            svgs[entry["name"]] = out_dir / f"{entry['file']}.svg"
            svgs[entry["name"]].write_text("<svg/>", encoding="utf-8")
            if entry["baked"]:
                grounds[entry["name"]] = out_dir / f"{entry['file']}.ground.svg"
                grounds[entry["name"]].write_text("<svg/>", encoding="utf-8")
        return {"svgs": svgs, "ground_svgs": grounds, "status": None}
    return render


def test_cache_hit_does_not_render_twice(decorated_template: Path, tmp_path: Path, monkeypatch) -> None:
    calls: list = []
    monkeypatch.setattr(decorations, "decoration_tools_status", lambda: None)
    monkeypatch.setattr(decorations, "render_decoration_svgs", _fake_render(calls))

    first = load_template_html(decorated_template, cache_root=tmp_path / "c", bundled_root=_no_bundle(tmp_path))
    second = load_template_html(decorated_template, cache_root=tmp_path / "c", bundled_root=_no_bundle(tmp_path))

    assert len(calls) == 1
    assert first.dir == second.dir == tmp_path / "c" / cache_key(decorated_template)
    assert first.decorations_status is None
    assert len(second.svg_paths) == 11
    assert set(second.ground_svg_paths) == {"Section Header"}
    assert (second.dir / "theme.css").read_text(encoding="utf-8") == second.theme_css
    assert json.loads((second.dir / "template.json").read_text(encoding="utf-8"))["sourceHash"] == second.dir.name


def test_failed_render_is_not_cached_as_success(decorated_template: Path, tmp_path: Path, monkeypatch) -> None:
    calls: list = []
    failed = {"status": "unavailable", "missing": [], "install": "", "note": "boom"}

    def render(*args, **kwargs):
        calls.append(1)
        return {"svgs": {}, "ground_svgs": {}, "status": failed}

    monkeypatch.setattr(decorations, "decoration_tools_status", lambda: None)
    monkeypatch.setattr(decorations, "render_decoration_svgs", render)
    for _ in range(2):
        result = load_template_html(decorated_template, cache_root=tmp_path / "c", bundled_root=_no_bundle(tmp_path))
        assert result.decorations_status == failed
        assert result.svg_paths == {}
    assert len(calls) == 1  # memoised for this process, never written as a success


def test_cache_key_changes_with_converter_version(decorated_template: Path, monkeypatch) -> None:
    before = cache_key(decorated_template)
    monkeypatch.setattr(template, "CONVERTER_VERSION", "test-bump")
    assert cache_key(decorated_template) != before
    assert len(before) == 16


def test_missing_tools_still_converts_structure(decorated_template: Path, tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr(decorations, "soffice_path", lambda: None)
    for _ in range(2):
        result = load_template_html(decorated_template, cache_root=tmp_path / "c", bundled_root=_no_bundle(tmp_path))
        status = result.decorations_status
        assert status is not None and status["status"] == "unavailable"
        assert "libreoffice" in status["missing"]
        assert set(status) == {"status", "missing", "install", "note"}
        assert "decorations" in status["note"]
        assert result.svg_paths == {}
        assert ":root{" in result.theme_css
        assert (result.dir / "template.json").is_file()


def test_bundled_artifacts_used_without_libreoffice(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr(decorations, "soffice_path", lambda: None)
    result = load_template_html(TEMPLATES_DIR / "blank-light.pptx", cache_root=tmp_path / "c")
    assert result.dir == TEMPLATES_DIR / "html" / "blank-light"
    assert result.decorations_status is None
    assert len(result.svg_paths) == len(result.info["layouts"])
    assert not (tmp_path / "c").exists()


# ── LibreOffice render ───────────────────────────────────────────────────

@_needs_lo
def test_render_strips_sentinel_from_band_layout(decorated_template: Path, tmp_path: Path) -> None:
    result = load_template_html(decorated_template, cache_root=tmp_path / "c", bundled_root=_no_bundle(tmp_path))
    assert result.decorations_status is None
    svg = result.svg_paths["Title Only"].read_text(encoding="utf-8")
    assert "<path" in svg or "<image" in svg
    assert "0.392" not in svg and "#010203" not in svg.lower()
    fills = [tuple(float(c) * 255 / 100 for c in m)
             for m in re.findall(r"rgb\(([\d.]+)%, ([\d.]+)%, ([\d.]+)%\)", svg)]
    assert any(all(abs(c - t) <= 1 for c, t in zip(fill, DECORATED_BAND_RGB)) for fill in fills)
    assert "<image" in svg  # master logo
    assert "Section Header" in result.ground_svg_paths
    blank = result.svg_paths["Blank"].read_text(encoding="utf-8")
    assert "<image" not in blank  # showMasterSp="0"


# ── bundled artifacts (drift guard) ──────────────────────────────────────

@pytest.mark.parametrize("stem", _BUNDLED)
def test_bundled_artifacts_match_fresh_conversion(stem: str, tmp_path: Path) -> None:
    pptx = TEMPLATES_DIR / f"{stem}.pptx"
    bundled = TEMPLATES_DIR / "html" / stem
    committed_json = (bundled / "template.json").read_text(encoding="utf-8")
    info = json.loads(committed_json)
    assert info["sourceHash"] == cache_key(pptx), (
        "bundled HTML artifacts are stale: run `uv run python scripts/build_bundled_template_html.py`")

    fresh = load_template_html(pptx, decorations=False, cache_root=tmp_path / "c", bundled_root=_no_bundle(tmp_path))
    assert fresh.theme_css == (bundled / "theme.css").read_text(encoding="utf-8")
    assert template_json_text(fresh.info) == committed_json

    for layout in info["layouts"]:
        assert (bundled / "layouts" / f"{layout['file']}.svg").is_file()
        if layout.get("groundSvg"):
            assert (bundled / "layouts" / f"{layout['file']}.ground.svg").is_file()

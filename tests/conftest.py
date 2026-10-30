# Copyright Amazon.com, Inc. or its affiliates. All Rights Reserved.
# SPDX-License-Identifier: MIT-0
"""Test configuration — add servers/remote/ and sdpm/ to sys.path."""

import sys
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Emu

_root = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(_root))
sys.path.insert(0, str(_root / "servers" / "remote"))
sys.path.insert(0, str(_root / "sdpm"))

# Standard slide dimensions in EMU
_W_16X9 = 12192000
_H_16X9 = 6858000
_W_4X3 = 9144000
_H_4X3 = 6858000


@pytest.fixture(autouse=True)
def _no_js_check(monkeypatch):
    """No test launches a browser for the HTML JS check unless it sets SDPM_JS_CHECK itself."""
    monkeypatch.setenv("SDPM_JS_CHECK", "0")


@pytest.fixture(autouse=True)
def _isolated_template_html_cache(tmp_path_factory, monkeypatch):
    """HTML template conversions go to a per-session temp cache, not the checkout's .cache."""
    import sdpm.config

    monkeypatch.setattr(sdpm.config, "CACHE_DIR", tmp_path_factory.getbasetemp() / "sdpm-cache")


@pytest.fixture()
def template_16x9() -> Path:
    """Path to the bundled 16:9 blank-dark template (no copy needed)."""
    return _root / "sdpm" / "templates" / "blank-dark.pptx"


@pytest.fixture()
def template_4x3(tmp_path: Path) -> Path:
    """Generate a 4:3 template from blank-dark.pptx by resizing slide dimensions.

    The roundtrip depends on layout name "Blank" existing in the template,
    so we cannot use a bare Presentation() — we must base it on blank-dark.pptx.
    """
    src = _root / "sdpm" / "templates" / "blank-dark.pptx"
    prs = Presentation(str(src))
    prs.slide_width = Emu(_W_4X3)
    prs.slide_height = Emu(_H_4X3)
    out = tmp_path / "blank-dark-4x3.pptx"
    prs.save(str(out))
    return out


# Geometry of the decorated_template fixture, in px on the 1920-wide canvas.
DECORATED_LOGO_PX = (1800, 1020, 64, 32)
DECORATED_COPYRIGHT_PX = (40, 1040, 400, 30)
DECORATED_BAND_RGB = (0xFF, 0x99, 0x00)


@pytest.fixture(scope="session")
def decorated_template(tmp_path_factory) -> Path:
    """blank-light.pptx with template decorations, built with python-pptx.

    - slide master 0: a 64x32 PNG logo and a copyright text box
    - layout "Title Only": a full-width header band containing its title placeholder
    - layout "Blank": showMasterSp="0" (hides the master logo / copyright)
    - layout "Section Header": a gradient background (non-solid ground)
    """
    import io

    from lxml import etree
    from PIL import Image
    from pptx.dml.color import RGBColor
    from pptx.shapes.autoshape import Shape

    prs = Presentation(str(_root / "sdpm" / "templates" / "blank-light.pptx"))
    emu = prs.slide_width / 1920
    master = prs.slide_masters[0]
    layouts = {layout.name: layout for layout in master.slide_layouts}

    def px(*values):
        return [int(round(v * emu)) for v in values]

    png = io.BytesIO()
    Image.new("RGB", (64, 32), (200, 30, 30)).save(png, format="PNG")
    png.seek(0)
    _, r_id = master.part.get_or_add_image_part(png)
    tree = master.shapes._spTree
    tree.add_pic(master.shapes._next_shape_id, "Logo", "", r_id, *px(*DECORATED_LOGO_PX))
    copyright_sp = tree.add_textbox(master.shapes._next_shape_id, "Copyright", *px(*DECORATED_COPYRIGHT_PX))
    Shape(copyright_sp, master.shapes).text_frame.text = "© Example Corp"

    title_only = layouts["Title Only"]
    title = next(ph for ph in title_only.placeholders if ph.placeholder_format.idx == 0)
    top = max(0, title.top - int(10 * emu))
    band_sp = title_only.shapes._spTree.add_autoshape(
        title_only.shapes._next_shape_id, "Header Band", "rect",
        0, top, prs.slide_width, title.height + int(20 * emu),
    )
    band = Shape(band_sp, title_only.shapes)
    band.fill.solid()
    band.fill.fore_color.rgb = RGBColor(*DECORATED_BAND_RGB)
    band.line.fill.background()

    layouts["Blank"]._element.set("showMasterSp", "0")

    ns = "http://schemas.openxmlformats.org/presentationml/2006/main"
    a = "http://schemas.openxmlformats.org/drawingml/2006/main"
    c_sld = layouts["Section Header"]._element.find(f"{{{ns}}}cSld")
    bg = etree.fromstring(
        f'<p:bg xmlns:p="{ns}" xmlns:a="{a}"><p:bgPr><a:gradFill><a:gsLst>'
        '<a:gs pos="0"><a:srgbClr val="112233"/></a:gs>'
        '<a:gs pos="100000"><a:srgbClr val="445566"/></a:gs>'
        "</a:gsLst><a:lin ang=\"0\" scaled=\"0\"/></a:gradFill><a:effectLst/></p:bgPr></p:bg>"
    )
    c_sld.insert(0, bg)

    out = tmp_path_factory.mktemp("decorated") / "decorated.pptx"
    prs.save(str(out))
    return out
